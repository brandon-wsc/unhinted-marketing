"""Parse pdf/docx/md/txt attachments into ordered document blocks (ADR 0047 §4)."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field


class DocumentParseError(Exception):
    """Raised when an attachment cannot be parsed into text blocks."""


class NoExtractableTextError(DocumentParseError):
    """ADR 0048 §1 — the document parsed cleanly but carries no text layer
    (scanned / image-only). ``page_count`` lets the ingest job park the row
    at ``needs_ocr`` with enough profile for a client to drive OCR."""

    def __init__(self, message: str, *, page_count: int = 0) -> None:
        super().__init__(message)
        self.page_count = page_count


@dataclass
class TranscriptPage:
    """One page of client-supplied OCR text (ADR 0048 §2)."""

    page: int
    text: str


@dataclass
class DocumentBlock:
    kind: str  # "heading" | "paragraph" | "table"
    text: str
    level: int | None = None
    page: int | None = None  # 1-based for PDF sources, None otherwise
    rows: list[list[str]] | None = None  # table cells; rows[0] is the header


@dataclass
class ParsedDocument:
    blocks: list[DocumentBlock] = field(default_factory=list)


_MD_HEADING_RE = re.compile(r"^(#{1,6})(?:\s+(.*?))?\s*$")
_TABLE_SEPARATOR_CELL_RE = re.compile(r"^:?-+:?$")
_DOCX_HEADING_STYLE_RE = re.compile(r"^Heading (\d+)$", re.IGNORECASE)
_BLANK_RUN_RE = re.compile(r"\n\s*\n")


def _render_table_row(cells: list[str]) -> str:
    return "| " + " | ".join(cells) + " |"


def _render_table_text(rows: list[list[str]]) -> str:
    return "\n".join(_render_table_row(row) for row in rows)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("File is not valid UTF-8 text") from exc


def _parse_txt(text: str) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    for para in _BLANK_RUN_RE.split(text):
        stripped = para.strip()
        if stripped:
            blocks.append(DocumentBlock(kind="paragraph", text=stripped))
    return blocks


def _table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator_row(cells: list[str]) -> bool:
    filled = [c for c in cells if c]
    return bool(filled) and all(_TABLE_SEPARATOR_CELL_RE.match(c) for c in filled)


def _parse_markdown(text: str) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    para_lines: list[str] = []
    table_lines: list[str] = []

    def flush_paragraph() -> None:
        if para_lines:
            para_text = "\n".join(para_lines).strip()
            if para_text:
                blocks.append(DocumentBlock(kind="paragraph", text=para_text))
            para_lines.clear()

    def flush_table() -> None:
        if table_lines:
            rows = []
            for line in table_lines:
                cells = _table_cells(line)
                if not _is_separator_row(cells):
                    rows.append(cells)
            blocks.append(DocumentBlock(kind="table", text=_render_table_text(rows), rows=rows))
            table_lines.clear()

    for line in text.split("\n"):
        stripped = line.strip()
        heading = _MD_HEADING_RE.match(stripped)
        if heading:
            flush_paragraph()
            flush_table()
            blocks.append(
                DocumentBlock(
                    kind="heading",
                    text=(heading.group(2) or "").strip(),
                    level=len(heading.group(1)),
                )
            )
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_lines.append(line)
            continue
        flush_table()
        if not stripped:
            flush_paragraph()
            continue
        para_lines.append(line)
    flush_paragraph()
    flush_table()
    return blocks


def _parse_docx(data: bytes) -> list[DocumentBlock]:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError) as exc:
        raise DocumentParseError("Could not read DOCX file") from exc

    blocks: list[DocumentBlock] = []
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = _DOCX_HEADING_STYLE_RE.match(item.style.name or "")
            if style:
                blocks.append(DocumentBlock(kind="heading", text=text, level=int(style.group(1))))
            else:
                blocks.append(DocumentBlock(kind="paragraph", text=text))
        elif isinstance(item, Table):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            blocks.append(DocumentBlock(kind="table", text=_render_table_text(rows), rows=rows))
    return blocks


def _reflow_pdf_text(text: str) -> str:
    """Reflow pdfium's visual line-wrap output into paragraph prose.

    ``get_text_range`` breaks at every rendered line (column width), not at
    sentence or paragraph ends — feeding it raw would make the chunker emit
    one fragment per line. ``￾`` is pdfium's soft-break marker (word split
    at the wrap): it joins with no space. A hyphen at a wrap joins without
    adding space, keeping the hyphen. Other single newlines become a space;
    blank-line runs (when a producer emits them) stay paragraph breaks.
    """
    text = text.replace("\r\n", "\n")
    text = text.replace("￾\n", "").replace("￾", "")
    text = re.sub(r"-\n(?=\S)", "-", text)
    text = _BLANK_RUN_RE.sub("\x00", text)
    return text.replace("\n", " ").replace("\x00", "\n\n")


def _parse_pdf(data: bytes) -> list[DocumentBlock]:
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise DocumentParseError("Could not read PDF file") from exc

    n_pages = len(doc)
    outline: dict[int, list[DocumentBlock]] = {}
    for bookmark in doc.get_toc():
        title = (bookmark.get_title() or "").strip()
        dest = bookmark.get_dest()
        if not title or dest is None:
            continue
        index = dest.get_index()
        if index < 0:
            continue
        outline.setdefault(index, []).append(
            DocumentBlock(
                kind="heading",
                text=title,
                level=bookmark.level + 1,
                page=index + 1,
            )
        )

    blocks: list[DocumentBlock] = []
    for i in range(n_pages):
        blocks.extend(outline.pop(i, []))
        page = doc[i]
        textpage = page.get_textpage()
        try:
            text = _reflow_pdf_text(textpage.get_text_range() or "")
        finally:
            textpage.close()
            page.close()
        for para in _BLANK_RUN_RE.split(text):
            stripped = para.strip()
            if stripped:
                blocks.append(DocumentBlock(kind="paragraph", text=stripped, page=i + 1))
    for index in sorted(outline):
        blocks.extend(outline[index])
    doc.close()

    if not any(block.kind == "paragraph" for block in blocks):
        raise NoExtractableTextError(
            "PDF has no extractable text — likely a scanned or image-only file",
            page_count=n_pages,
        )
    return blocks


def parse_transcript(pages: list[TranscriptPage]) -> ParsedDocument:
    """ADR 0048 §3 — client OCR output → the same ``ParsedDocument`` shape
    ``parse_document`` returns, so chunking/profiling run unchanged. One
    paragraph block per blank-line run; blank pages contribute nothing."""
    blocks: list[DocumentBlock] = []
    for entry in pages:
        for para in _BLANK_RUN_RE.split(entry.text):
            stripped = para.strip()
            if stripped:
                blocks.append(
                    DocumentBlock(kind="paragraph", text=stripped, page=entry.page)
                )
    if not blocks:
        raise DocumentParseError("Transcript contains no readable text")
    return ParsedDocument(blocks=blocks)


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return ParsedDocument(blocks=_parse_pdf(data))
    if lower.endswith(".docx"):
        return ParsedDocument(blocks=_parse_docx(data))
    if lower.endswith((".md", ".markdown")):
        return ParsedDocument(blocks=_parse_markdown(_decode_text(data)))
    if lower.endswith(".txt"):
        return ParsedDocument(blocks=_parse_txt(_decode_text(data)))
    raise DocumentParseError(f"Unsupported file type: {filename}")

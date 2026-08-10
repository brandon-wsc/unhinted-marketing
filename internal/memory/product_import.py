"""Parse flexible CSV / Excel product imports into row dicts."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

SKU_KEYS = ("sku", "product_sku", "product_code", "code", "id", "貨號", "編號", "sku編號")
NAME_KEYS = ("name", "product_name", "product", "title", "名稱", "產品", "產品名稱")
MAX_IMPORT_COLUMNS = 50


def _usable_headers(headers: list[Any]) -> list[str]:
    out: list[str] = []
    for h in headers:
        if h is None:
            continue
        text = str(h).strip()
        if text:
            out.append(text)
    return out


def _column_limit_error(headers: list[Any]) -> str | None:
    count = len(_usable_headers(headers))
    if count <= MAX_IMPORT_COLUMNS:
        return None
    return (
        f"Too many columns ({count}; max {MAX_IMPORT_COLUMNS}). "
        "Remove unused columns and keep product code, name, and selling details."
    )


def _norm_header(raw: str) -> str:
    return re.sub(r"\s+", " ", raw.strip().lower())


def _pick_key(headers: list[str], candidates: tuple[str, ...]) -> str | None:
    normalized = {_norm_header(h): h for h in headers}
    for cand in candidates:
        if cand.lower() in normalized:
            return normalized[cand.lower()]
    return None


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def build_search_document(profile: dict[str, str], *, sku: str, name: str) -> str:
    parts: list[str] = []
    for key in (name, sku):
        if key and key not in parts:
            parts.append(key)
    for value in profile.values():
        text = value.strip()
        if text and text not in parts:
            parts.append(text)
    return " | ".join(parts)


def rows_from_dicts(raw_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    """Return (parsed_rows, row_errors). Each parsed row: sku, name, profile, search_document."""
    if not raw_rows:
        return [], ["File has no data rows"]

    headers = list(raw_rows[0].keys())
    limit_err = _column_limit_error(headers)
    if limit_err:
        return [], [limit_err]

    sku_col = _pick_key(headers, SKU_KEYS)
    name_col = _pick_key(headers, NAME_KEYS)

    parsed: list[dict[str, Any]] = []
    errors: list[str] = []
    for idx, raw in enumerate(raw_rows, start=2):  # 1-based + header
        profile = {
            str(k): _cell_str(v)
            for k, v in raw.items()
            if k is not None and _cell_str(v)
        }
        if not profile:
            errors.append(f"Row {idx}: empty")
            continue

        sku = _cell_str(raw.get(sku_col)) if sku_col else ""
        name = _cell_str(raw.get(name_col)) if name_col else ""
        if not sku and name:
            sku = name
        if not name and sku:
            name = sku
        if not sku:
            errors.append(f"Row {idx}: missing product code / name")
            continue
        if len(sku) > 200:
            errors.append(f"Row {idx}: product code too long")
            continue
        if len(name) > 500:
            name = name[:500]

        parsed.append(
            {
                "sku": sku,
                "name": name,
                "profile": profile,
                "search_document": build_search_document(profile, sku=sku, name=name),
            }
        )
    return parsed, errors


def parse_csv_bytes(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    text = data.decode("utf-8-sig")
    if not text.strip():
        return [], ["File is empty"]
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return [], ["CSV has no header row"]
    limit_err = _column_limit_error(list(reader.fieldnames))
    if limit_err:
        return [], [limit_err]
    raw_rows = [dict(row) for row in reader]
    return rows_from_dicts(raw_rows)


def parse_xlsx_bytes(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for Excel import") from exc

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return [], ["Workbook has no sheet"]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["File is empty"]
    headers = [_cell_str(h) or f"col_{i}" for i, h in enumerate(header_row)]
    if not any(headers):
        return [], ["Excel has no header row"]
    limit_err = _column_limit_error(headers)
    if limit_err:
        return [], [limit_err]

    raw_rows: list[dict[str, Any]] = []
    for values in rows_iter:
        if values is None or all(v is None or _cell_str(v) == "" for v in values):
            continue
        raw_rows.append(
            {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        )
    return rows_from_dicts(raw_rows)


def parse_product_upload(filename: str, data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv_bytes(data)
    if lower.endswith(".xlsx"):
        return parse_xlsx_bytes(data)
    return [], ["Unsupported file type — use .csv or .xlsx"]

"""Unit tests for ADR 0047 document parsing (pdf/docx/md/txt -> ordered blocks).

Production module under test: internal/memory/document_parse.py (not yet written).
Imports stay inside helpers/test bodies so collection succeeds while the module
is missing — every test should fail, not error at import time.
"""

from __future__ import annotations

import io

import pytest


def _parse(filename: str, data: bytes):
    from internal.memory.document_parse import parse_document

    return parse_document(filename, data)


def _parse_error():
    from internal.memory.document_parse import DocumentParseError

    return DocumentParseError


def _pdf_stream(text: str) -> bytes:
    body = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    return b"<< /Length " + str(len(body)).encode() + b" >>\nstream\n" + body + b"\nendstream"


def _pdf_doc(objs: list[bytes]) -> bytes:
    out = b"%PDF-1.4\n"
    for i, obj in enumerate(objs, 1):
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    trailer = b"trailer\n<< /Root 1 0 R /Size " + str(len(objs) + 1).encode() + b" >>\n%%EOF"
    return out + trailer


def _make_pdf(pages: list[str | None]) -> bytes:
    """Minimal PDF (no xref — pdfium reconstructs). None = page without /Contents."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        contents = f"/Contents {4 + 2 * i} 0 R" if text is not None else ""
        objs.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> {contents} >>"
            ).encode()
        )
        objs.append(_pdf_stream(text) if text is not None else b"<< >>")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    return _pdf_doc(objs)


def _make_pdf_with_outline() -> bytes:
    return _pdf_doc(
        [
            b"<< /Type /Catalog /Pages 2 0 R /Outlines 3 0 R >>",
            b"<< /Type /Pages /Kids [4 0 R 6 0 R] /Count 2 >>",
            b"<< /Type /Outlines /First 8 0 R /Last 9 0 R /Count 2 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 10 0 R >> >> /Contents 5 0 R >>",
            _pdf_stream("Page one body"),
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 10 0 R >> >> /Contents 7 0 R >>",
            _pdf_stream("Page two body"),
            b"<< /Title (Intro) /Parent 3 0 R /Next 9 0 R /Dest [4 0 R /Fit] >>",
            b"<< /Title (Details) /Parent 3 0 R /Prev 8 0 R /Dest [6 0 R /Fit] >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
    )


def _docx_bytes(build) -> bytes:
    from docx import Document

    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_markdown_headings_carry_levels_and_order() -> None:
    md = "# 產品目錄\n\n第一段說明。\n\n## 優惠詳情\n\n第二段內容。\n"
    doc = _parse("catalog.md", md.encode())
    blocks = list(doc.blocks)
    assert [b.kind for b in blocks] == ["heading", "paragraph", "heading", "paragraph"]
    assert blocks[0].text == "產品目錄" and blocks[0].level == 1
    assert blocks[2].text == "優惠詳情" and blocks[2].level == 2
    assert "第一段" in blocks[1].text and "第二段" in blocks[3].text


def test_parse_markdown_pipe_table_stays_single_block() -> None:
    md = "## 價目表\n\n| 名稱 | 價格 |\n| --- | --- |\n| 奶茶 | 48 |\n| 咖啡 | 52 |\n\n表後文字。\n"
    doc = _parse("prices.md", md.encode())
    tables = [b for b in doc.blocks if b.kind == "table"]
    assert len(tables) == 1
    table = tables[0]
    assert table.rows is not None
    assert table.rows[0][0] == "名稱" and table.rows[0][1] == "價格"
    assert len(table.rows) == 3
    assert "奶茶" in table.text and "咖啡" in table.text
    assert any(b.kind == "paragraph" and "表後文字" in b.text for b in doc.blocks)


def test_parse_txt_gives_paragraph_blocks() -> None:
    doc = _parse("notes.txt", "第一段內容。\n\n第二段內容。\n".encode())
    blocks = [b for b in doc.blocks if b.text.strip()]
    assert [b.kind for b in blocks] == ["paragraph", "paragraph"]
    assert all(b.level is None for b in blocks)
    assert "第一段" in blocks[0].text and "第二段" in blocks[1].text


def test_parse_docx_heading_styles_become_heading_blocks() -> None:
    def build(doc) -> None:
        doc.add_heading("產品目錄", level=1)
        doc.add_paragraph("第一段內容")
        doc.add_heading("優惠詳情", level=2)
        doc.add_paragraph("第二段內容")

    doc = _parse("brief.docx", _docx_bytes(build))
    blocks = list(doc.blocks)
    headings = [b for b in blocks if b.kind == "heading"]
    assert [h.text for h in headings] == ["產品目錄", "優惠詳情"]
    assert [h.level for h in headings] == [1, 2]
    paras = [b for b in blocks if b.kind == "paragraph"]
    assert "第一段內容" in paras[0].text and "第二段內容" in paras[1].text
    assert blocks.index(headings[0]) < blocks.index(paras[0]) < blocks.index(headings[1])


def test_parse_docx_table_is_atomic_block() -> None:
    def build(doc) -> None:
        doc.add_paragraph("前言")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "名稱"
        table.rows[0].cells[1].text = "價格"
        table.rows[1].cells[0].text = "奶茶"
        table.rows[1].cells[1].text = "48"
        doc.add_paragraph("後記")

    doc = _parse("list.docx", _docx_bytes(build))
    tables = [b for b in doc.blocks if b.kind == "table"]
    assert len(tables) == 1
    table = tables[0]
    assert table.rows is not None and len(table.rows) == 2
    assert table.rows[0][0] == "名稱" and table.rows[1][0] == "奶茶"
    assert any(b.kind == "paragraph" and "後記" in b.text for b in doc.blocks)


def test_parse_pdf_extracts_text_per_page_with_page_numbers() -> None:
    doc = _parse("sheet.pdf", _make_pdf(["Page one body", "Page two body"]))
    paras = [b for b in doc.blocks if b.kind == "paragraph" and b.text.strip()]
    assert [b.page for b in paras] == [1, 2]
    assert "Page one" in paras[0].text and "Page two" in paras[1].text


def test_parse_pdf_outline_becomes_heading_blocks() -> None:
    doc = _parse("outlined.pdf", _make_pdf_with_outline())
    blocks = list(doc.blocks)
    headings = [b for b in blocks if b.kind == "heading"]
    assert [h.text for h in headings] == ["Intro", "Details"]
    assert all(h.level == 1 for h in headings)
    assert headings[0].page == 1 and headings[1].page == 2
    paras = [b for b in blocks if b.kind == "paragraph" and b.text.strip()]
    assert blocks.index(headings[0]) < blocks.index(paras[0]) < blocks.index(headings[1])


def test_parse_pdf_reflows_column_wrapped_lines_into_prose() -> None:
    """Pdfium emits \\n at every visual line wrap — wrapped lines must rejoin
    into flowing sentences or the chunker cuts one chunk per line (ADR 0047
    §4: chunks are readable slices, not layout artifacts)."""
    ops = (
        b"BT /F1 12 Tf 50 700 Td (First wrapped line of a sentence still) Tj ET "
        b"BT /F1 12 Tf 50 686 Td (continuing on the next row.) Tj ET"
    )
    body = b"<< /Length " + str(len(ops)).encode() + b" >>\nstream\n" + ops + b"\nendstream"
    pdf = _pdf_doc(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            body,
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
    )
    doc = _parse("wrap.pdf", pdf)
    paras = [b for b in doc.blocks if b.kind == "paragraph"]
    assert len(paras) == 1
    assert "\n" not in paras[0].text
    assert "sentence still continuing on the next row." in paras[0].text


def test_reflow_pdf_text_joins_soft_breaks_and_hyphen_wraps() -> None:
    from internal.memory.document_parse import _reflow_pdf_text

    # pdfium soft-break marker \ufffe joins with no space ("com￾\npile")
    assert _reflow_pdf_text("com￾\npile is done.") == "compile is done."
    # A wrapped hyphen keeps the hyphen ("well-\nknown")
    assert _reflow_pdf_text("well-\nknown trick") == "well-known trick"
    # Plain wraps become a space; blank-line runs stay paragraph breaks
    assert _reflow_pdf_text("line one\nline two") == "line one line two"
    assert _reflow_pdf_text("para one\n\npara two") == "para one\n\npara two"


def test_parse_scanned_pdf_raises_readable_error() -> None:
    with pytest.raises(_parse_error()):
        _parse("scan.pdf", _make_pdf([None]))


def test_parse_scanned_pdf_raises_no_extractable_text_with_page_count() -> None:
    """ADR 0048 §1: a no-text-layer PDF raises ``NoExtractableTextError`` — a
    ``DocumentParseError`` subclass carrying ``page_count`` so the ingest job
    can park the row at ``needs_ocr`` with a real page count for the UI
    instead of failing it blind."""
    from internal.memory.document_parse import NoExtractableTextError

    with pytest.raises(NoExtractableTextError) as excinfo:
        _parse("scan.pdf", _make_pdf([None, None, None]))
    assert isinstance(excinfo.value, _parse_error())
    assert excinfo.value.page_count == 3


def test_parse_transcript_pages_become_page_numbered_paragraphs() -> None:
    """ADR 0048 §3 — ``parse_transcript(pages)`` turns client OCR output into
    the same ``ParsedDocument`` shape ``parse_document`` returns: one
    paragraph block per blank-line run with ``page`` carried through, and
    blank pages contribute nothing."""
    from internal.memory.document_parse import TranscriptPage, parse_transcript

    doc = parse_transcript(
        [
            TranscriptPage(page=1, text="第一頁第一段。\n\n第一頁第二段。"),
            TranscriptPage(page=2, text="   "),
            TranscriptPage(page=3, text="第三頁內容。"),
        ]
    )
    paras = [b for b in doc.blocks if b.kind == "paragraph"]
    assert [b.page for b in paras] == [1, 1, 3]
    assert "第一頁第一段" in paras[0].text
    assert "第一頁第二段" in paras[1].text
    assert "第三頁內容" in paras[2].text


def test_parse_transcript_all_blank_raises_readable_error() -> None:
    """A transcript that contributes no text is a parse failure, not a false
    ``ready`` (ADR 0048 §9 — same honesty posture as ADR 0047 §10)."""
    from internal.memory.document_parse import TranscriptPage, parse_transcript

    with pytest.raises(_parse_error()):
        parse_transcript([TranscriptPage(page=1, text="  \n\n  ")])


def test_parse_corrupt_pdf_raises_readable_error() -> None:
    with pytest.raises(_parse_error()):
        _parse("broken.pdf", b"this is not a pdf document")


def test_parse_unsupported_extension_raises_readable_error() -> None:
    with pytest.raises(_parse_error()):
        _parse("data.bin", b"\x00\x01\x02\x03")
